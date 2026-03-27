import json

from openpyxl import Workbook
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


class ExportService:
    # Backward-compatible mapping fields used by import flow.
    GENERAL_FIELDS = [
        ('id', 'Tool ID'),
        ('tool_head', 'Tool Head'),
        ('tool_type', 'Tool type'),
        ('description', 'Description'),
        ('geom_x', 'Geom X'),
        ('geom_z', 'Geom Z'),
        ('radius', 'Radius'),
        ('nose_corner_radius', 'Nose R / Corner R'),
        ('holder_code', 'Holder code'),
        ('holder_link', 'Holder link'),
        ('holder_add_element', 'Add. Element'),
        ('holder_add_element_link', 'Add. Element link'),
        ('cutting_type', 'Cutting component type'),
        ('cutting_code', 'Cutting component code'),
        ('cutting_link', 'Cutting component link'),
        ('cutting_add_element', 'Add. Insert/Drill/Mill'),
        ('cutting_add_element_link', 'Add. Insert/Drill/Mill link'),
        ('drill_nose_angle', 'Nose angle'),
        ('mill_cutting_edges', 'Cutting edges'),
        ('notes', 'Notes'),
    ]

    # V2 export keeps stable tool metadata columns and expands components dynamically.
    EXPORT_BASE_FIELDS = [
        ('export_format', 'Export format'),
        ('id', 'Tool ID'),
        ('tool_head', 'Tool Head'),
        ('tool_type', 'Tool type'),
        ('description', 'Description'),
        ('geom_x', 'Geom X'),
        ('geom_z', 'Geom Z'),
        ('radius', 'Radius'),
        ('nose_corner_radius', 'Nose R / Corner R'),
        ('cutting_type', 'Cutting component type'),
        ('drill_nose_angle', 'Nose angle'),
        ('mill_cutting_edges', 'Cutting edges'),
        ('notes', 'Notes'),
    ]

    _ROLE_ORDER = ('holder', 'cutting', 'support')
    _ROLE_LABEL = {
        'holder': 'Holder',
        'cutting': 'Cutting',
        'support': 'Support',
    }
    _COMPONENT_SLOT_CAP = 3

    _TOOLTYPE_ROW_COLORS = [
        'EAF4FF',  # light blue
        'EDF8EA',  # light green
        'FFF4E8',  # light orange
        'F5ECFF',  # light violet
        'EAF8F8',  # light cyan
        'FFF9E3',  # light yellow
        'FDEEEF',  # light rose
        'ECEFF3',  # light slate
    ]

    IMPORT_DEFAULTS = {
        'id': '',
        'tool_head': 'HEAD1',
        'tool_type': 'O.D Turning',
        'description': '',
        'geom_x': 0.0,
        'geom_z': 0.0,
        'radius': 0.0,
        'nose_corner_radius': 0.0,
        'holder_code': '',
        'holder_link': '',
        'holder_add_element': '',
        'holder_add_element_link': '',
        'cutting_type': 'Insert',
        'cutting_code': '',
        'cutting_link': '',
        'cutting_add_element': '',
        'cutting_add_element_link': '',
        'notes': '',
        'drill_nose_angle': 0.0,
        'mill_cutting_edges': 0,
        'geometry_profiles': [],
        'support_parts': [],
        'stl_path': '',
    }

    def _normalize_number(self, value):
        if value is None or value == '':
            return None
        try:
            return float(value)
        except Exception:
            return value

    @staticmethod
    def _normalize_tool_head(value) -> str:
        head = str(value or 'HEAD1').strip().upper()
        return head if head in {'HEAD1', 'HEAD2'} else 'HEAD1'

    @staticmethod
    def _parse_json_list_or_empty(value):
        if isinstance(value, list):
            return value
        if value is None:
            return []
        text = str(value).strip()
        if not text:
            return []
        import json
        try:
            parsed = json.loads(text)
            return parsed if isinstance(parsed, list) else []
        except Exception:
            return []

    def _normalize_component_items(self, tool: dict) -> list[dict]:
        items = self._parse_json_list_or_empty(tool.get('component_items', []))

        normalized = []
        for idx, item in enumerate(items):
            if not isinstance(item, dict):
                continue
            role = str(item.get('role') or '').strip().lower()
            code = str(item.get('code') or '').strip()
            if role not in self._ROLE_ORDER or not code:
                continue
            label = str(item.get('label') or '').strip() or self._ROLE_LABEL[role]
            try:
                order = int(item.get('order', idx))
            except Exception:
                order = idx
            normalized.append(
                {
                    'role': role,
                    'label': label,
                    'code': code,
                    'link': str(item.get('link') or '').strip(),
                    'group': str(item.get('group') or '').strip(),
                    'order': order,
                }
            )

        if normalized:
            normalized.sort(key=lambda entry: int(entry.get('order', 0)))
            return normalized

        fallback = []

        def _add(role, label, code, link='', group=''):
            code_text = str(code or '').strip()
            if not code_text:
                return
            fallback.append(
                {
                    'role': role,
                    'label': str(label or '').strip() or self._ROLE_LABEL[role],
                    'code': code_text,
                    'link': str(link or '').strip(),
                    'group': str(group or '').strip(),
                    'order': len(fallback),
                }
            )

        cutting_type = str(tool.get('cutting_type', 'Insert') or 'Insert').strip() or 'Insert'
        _add('holder', 'Holder', tool.get('holder_code', ''), tool.get('holder_link', ''))
        _add('holder', 'Add. Element', tool.get('holder_add_element', ''), tool.get('holder_add_element_link', ''))
        _add('cutting', cutting_type, tool.get('cutting_code', ''), tool.get('cutting_link', ''))
        _add(
            'cutting',
            f'Add. {cutting_type}',
            tool.get('cutting_add_element', ''),
            tool.get('cutting_add_element_link', ''),
        )

        for part in self._parse_json_list_or_empty(tool.get('support_parts', [])):
            if isinstance(part, str):
                try:
                    part = json.loads(part)
                except Exception:
                    part = {'name': part, 'code': '', 'link': '', 'group': ''}
            if not isinstance(part, dict):
                continue
            _add(
                'support',
                str(part.get('name') or 'Part').strip() or 'Part',
                part.get('code', ''),
                part.get('link', ''),
                part.get('group', ''),
            )

        return fallback

    def _analyze_component_layout(self, tools: list[dict], slot_cap: int | None = None):
        cap = int(slot_cap or self._COMPONENT_SLOT_CAP)
        max_counts = {role: 0 for role in self._ROLE_ORDER}
        overflow = {role: False for role in self._ROLE_ORDER}

        for tool in tools or []:
            items = self._normalize_component_items(tool)
            by_role = {role: [] for role in self._ROLE_ORDER}
            for item in items:
                role = item.get('role')
                if role in by_role:
                    by_role[role].append(item)
            for role in self._ROLE_ORDER:
                count = len(by_role[role])
                max_counts[role] = max(max_counts[role], count)
                if count > cap:
                    overflow[role] = True

        slot_counts = {role: min(cap, max_counts[role]) for role in self._ROLE_ORDER}
        return {
            'slot_counts': slot_counts,
            'overflow': overflow,
            'slot_cap': cap,
        }

    def _component_columns(self, layout):
        columns = []
        for role in self._ROLE_ORDER:
            role_label = self._ROLE_LABEL[role]
            for slot in range(1, int(layout['slot_counts'][role]) + 1):
                columns.extend(
                    [
                        {
                            'key': f'{role}_{slot}_label',
                            'label': f'{role_label} {slot} label',
                            'role': role,
                            'slot': slot,
                            'attr': 'label',
                        },
                        {
                            'key': f'{role}_{slot}_code',
                            'label': f'{role_label} {slot} code',
                            'role': role,
                            'slot': slot,
                            'attr': 'code',
                        },
                        {
                            'key': f'{role}_{slot}_link',
                            'label': f'{role_label} {slot} link',
                            'role': role,
                            'slot': slot,
                            'attr': 'link',
                        },
                        {
                            'key': f'{role}_{slot}_group',
                            'label': f'{role_label} {slot} group',
                            'role': role,
                            'slot': slot,
                            'attr': 'group',
                        },
                    ]
                )

            if layout['overflow'][role]:
                columns.append(
                    {
                        'key': f'{role}_overflow_json',
                        'label': f'{role_label} overflow (JSON)',
                        'role': role,
                        'slot': None,
                        'attr': 'overflow',
                    }
                )
        return columns

    @staticmethod
    def _group_items_by_role(items: list[dict]):
        grouped = {'holder': [], 'cutting': [], 'support': []}
        for item in items or []:
            role = item.get('role')
            if role in grouped:
                grouped[role].append(item)
        for role in grouped:
            grouped[role].sort(key=lambda entry: int(entry.get('order', 0)))
        return grouped

    def read_excel_headers(self, filename: str) -> list[str]:
        wb = load_workbook(filename=filename, read_only=True, data_only=True)
        try:
            ws = wb.active
            first_row = next(ws.iter_rows(min_row=1, max_row=1, values_only=True), None)
            if not first_row:
                return []
            return [str(v).strip() for v in first_row if v is not None and str(v).strip()]
        finally:
            wb.close()

    def import_tools(self, filename: str, mapping: dict[str, str]) -> list[dict]:
        wb = load_workbook(filename=filename, read_only=True, data_only=True)
        try:
            ws = wb.active
            raw_headers = next(ws.iter_rows(min_row=1, max_row=1, values_only=True), None)
            if not raw_headers:
                return []
            headers = [str(v).strip() if v is not None else '' for v in raw_headers]
            header_to_idx = {h: i for i, h in enumerate(headers) if h}

            float_fields = {'geom_x', 'geom_z', 'radius', 'nose_corner_radius', 'drill_nose_angle'}
            int_fields = {'mill_cutting_edges'}
            list_fields = {'support_parts', 'geometry_profiles'}

            imported = []
            for row in ws.iter_rows(min_row=2, values_only=True):
                tool = dict(self.IMPORT_DEFAULTS)

                for field_key, excel_header in mapping.items():
                    if field_key not in self.IMPORT_DEFAULTS:
                        continue
                    idx = header_to_idx.get(excel_header)
                    if idx is None or idx >= len(row):
                        continue
                    raw = row[idx]
                    if raw is None:
                        continue

                    if field_key in float_fields:
                        try:
                            tool[field_key] = float(raw)
                        except Exception:
                            pass
                    elif field_key in int_fields:
                        try:
                            tool[field_key] = int(raw)
                        except Exception:
                            pass
                    elif field_key in list_fields:
                        tool[field_key] = self._parse_json_list_or_empty(raw)
                    else:
                        tool[field_key] = str(raw).strip()

                tool['tool_head'] = self._normalize_tool_head(tool.get('tool_head', 'HEAD1'))

                if not str(tool.get('id', '')).strip():
                    continue

                imported.append(tool)

            return imported
        finally:
            wb.close()

    def _write_tools_sheet(self, ws, tools: list[dict]):
        layout = self._analyze_component_layout(tools)
        component_columns = self._component_columns(layout)
        headers = [label for _key, label in self.EXPORT_BASE_FIELDS] + [col['label'] for col in component_columns]

        # Header row
        ws.append(headers)
        header_fill = PatternFill(fill_type='solid', fgColor='1F4E78')
        header_font = Font(color='FFFFFF', bold=True)
        thin = Side(style='thin', color='D0D7DE')
        border = Border(left=thin, right=thin, top=thin, bottom=thin)

        for col_idx in range(1, len(headers) + 1):
            cell = ws.cell(row=1, column=col_idx)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal='center', vertical='center')
            cell.border = border

        # Stable color mapping for each tool type present in export
        unique_tool_types = []
        for tool in tools or []:
            t = (tool.get('tool_type', '') or '').strip()
            if t and t not in unique_tool_types:
                unique_tool_types.append(t)
        tool_type_fill = {
            t: PatternFill(fill_type='solid', fgColor=self._TOOLTYPE_ROW_COLORS[idx % len(self._TOOLTYPE_ROW_COLORS)])
            for idx, t in enumerate(sorted(unique_tool_types))
        }

        numeric_keys = {'geom_x', 'geom_z', 'radius', 'nose_corner_radius', 'drill_nose_angle'}
        int_keys = {'mill_cutting_edges'}

        # Data rows
        for tool in tools or []:
            normalized_tool = dict(tool)
            normalized_tool['export_format'] = f"TOOL_EXPORT_V2;slot_cap={layout['slot_cap']}"

            row_values = []
            for key, _label in self.EXPORT_BASE_FIELDS:
                value = normalized_tool.get(key, '')
                if key in numeric_keys:
                    value = self._normalize_number(value)
                elif key in int_keys:
                    try:
                        value = int(value or 0)
                    except Exception:
                        value = value
                row_values.append(value)

            grouped_items = self._group_items_by_role(self._normalize_component_items(normalized_tool))
            for column in component_columns:
                role = column['role']
                attr = column['attr']
                if attr == 'overflow':
                    overflow_items = grouped_items[role][layout['slot_counts'][role]:]
                    row_values.append(json.dumps(overflow_items, ensure_ascii=False) if overflow_items else '')
                    continue

                slot_idx = int(column['slot']) - 1
                role_items = grouped_items[role]
                if slot_idx >= len(role_items):
                    row_values.append('')
                    continue
                row_values.append(role_items[slot_idx].get(attr, ''))

            ws.append(row_values)
            row_idx = ws.max_row

            fill = tool_type_fill.get((tool.get('tool_type', '') or '').strip())
            for col_idx in range(1, len(headers) + 1):
                cell = ws.cell(row=row_idx, column=col_idx)
                cell.border = border
                if fill is not None:
                    cell.fill = fill

                if col_idx <= len(self.EXPORT_BASE_FIELDS):
                    key = self.EXPORT_BASE_FIELDS[col_idx - 1][0]
                else:
                    key = ''
                if key in numeric_keys and isinstance(cell.value, (int, float)):
                    cell.number_format = '0.000'
                    cell.alignment = Alignment(horizontal='right', vertical='center')
                elif key in int_keys and isinstance(cell.value, int):
                    cell.number_format = '0'
                    cell.alignment = Alignment(horizontal='right', vertical='center')
                else:
                    cell.alignment = Alignment(horizontal='left', vertical='center')

        ws.freeze_panes = 'A2'
        ws.auto_filter.ref = ws.dimensions

        # Auto-size columns with sensible min/max limits.
        for col_idx in range(1, len(headers) + 1):
            letter = get_column_letter(col_idx)
            max_len = len(headers[col_idx - 1])
            for row_idx in range(2, ws.max_row + 1):
                val = ws.cell(row=row_idx, column=col_idx).value
                if val is None:
                    continue
                max_len = max(max_len, len(str(val)))
            ws.column_dimensions[letter].width = min(max(max_len + 2, 10), 52)

        # Slightly taller rows for readability.
        ws.row_dimensions[1].height = 24
        for row_idx in range(2, ws.max_row + 1):
            ws.row_dimensions[row_idx].height = 21

    def export_tools(self, filename: str, tools: list[dict]):
        wb = Workbook()
        head1_ws = wb.active
        head1_ws.title = 'HEAD1'
        head2_ws = wb.create_sheet('HEAD2')

        head1_tools = []
        head2_tools = []

        for tool in tools or []:
            normalized_head = self._normalize_tool_head(tool.get('tool_head', 'HEAD1'))
            tool_copy = dict(tool)
            tool_copy['tool_head'] = normalized_head
            if normalized_head == 'HEAD2':
                head2_tools.append(tool_copy)
            else:
                head1_tools.append(tool_copy)

        self._write_tools_sheet(head1_ws, head1_tools)
        self._write_tools_sheet(head2_ws, head2_tools)

        wb.save(filename)

    def export_tools_to_excel(self, conn, filename):
        # Backward-compatible path for legacy callers.
        rows = conn.execute('SELECT * FROM tools ORDER BY id').fetchall()
        tools = [dict(r) for r in rows]
        self.export_tools(filename, tools)
